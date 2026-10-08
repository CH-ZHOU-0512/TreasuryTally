"""Real in-memory DOCX with shared report content and embedded OFL font."""

from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from trust_receipt.reporting.graphics import font_bytes, graph_png
from trust_receipt.reporting.layout import report_sections, validate_export_view
from trust_receipt.reporting.models import BusinessReportView
from trust_receipt.reporting.renderer import EChartsRenderer, ExportUnavailable

FONT_NAME = "TreasuryTally Report Sans"


def _embed_font(payload: bytes) -> bytes:
    from lxml import etree as ET

    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    p = "http://schemas.openxmlformats.org/package/2006/relationships"
    key = uuid4()
    mask = bytes.fromhex(key.hex)[::-1]
    font = bytearray(font_bytes())
    for index in range(32):
        font[index] ^= mask[index % 16]
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as original, ZipFile(output, "w", ZIP_DEFLATED) as target:
        table = ET.fromstring(original.read("word/fontTable.xml"))
        node = ET.SubElement(table, f"{{{w}}}font", {f"{{{w}}}name": FONT_NAME})
        ET.SubElement(
            node,
            f"{{{w}}}embedRegular",
            {f"{{{r}}}id": "rIdReportFont", f"{{{w}}}fontKey": "{" + str(key).upper() + "}"},
        )
        relationships = (
            ET.fromstring(original.read("word/_rels/fontTable.xml.rels"))
            if "word/_rels/fontTable.xml.rels" in original.namelist()
            else ET.Element(f"{{{p}}}Relationships", nsmap={None: p})
        )
        ET.SubElement(
            relationships,
            f"{{{p}}}Relationship",
            {
                "Id": "rIdReportFont",
                "Type": r + "/font",
                "Target": "fonts/report.odttf",
            },
        )
        types = ET.fromstring(original.read("[Content_Types].xml"))
        ET.SubElement(
            types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            {
                "PartName": "/word/fonts/report.odttf",
                "ContentType": "application/vnd.openxmlformats-officedocument.obfuscatedFont",
            },
        )
        replacements = {
            "word/fontTable.xml": ET.tostring(table, encoding="utf-8", xml_declaration=True),
            "word/_rels/fontTable.xml.rels": ET.tostring(relationships, encoding="utf-8", xml_declaration=True),
            "[Content_Types].xml": ET.tostring(types, encoding="utf-8", xml_declaration=True),
        }
        for entry in original.infolist():
            if entry.filename not in replacements:
                target.writestr(entry, original.read(entry.filename))
        for name, contents in replacements.items():
            target.writestr(name, contents)
        target.writestr("word/fonts/report.odttf", bytes(font))
    return output.getvalue()


def export_docx(view: BusinessReportView, *, renderer: EChartsRenderer) -> bytes:
    with renderer.export_slot():
        return _export_docx(view, renderer=renderer)


def _export_docx(view: BusinessReportView, *, renderer: EChartsRenderer) -> bytes:
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as error:
        raise ExportUnavailable("EXPORT_UNAVAILABLE: Word export dependency missing") from error

    validate_export_view(view)
    doc = Document()
    doc.core_properties.author = "TreasuryTally"
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = view.title
    doc.core_properties.subject = ""
    doc.core_properties.comments = ""
    doc.core_properties.created = doc.core_properties.modified = datetime.now(UTC)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = FONT_NAME
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("asciiTheme", "eastAsiaTheme", "hAnsiTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)
        fonts.set(qn("w:eastAsia"), FONT_NAME)
        fonts.set(qn("w:cs"), FONT_NAME)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.add_paragraph(view.title, "Title")
    for content in report_sections(view):
        doc.add_heading(content.title, level=1)
        for index, text in enumerate(content.paragraphs):
            paragraph = doc.add_paragraph(text)
            if content.headers or (
                content.title == "完整账户索引" and len(content.paragraphs) <= 8 and index < len(content.paragraphs) - 1
            ):
                paragraph.paragraph_format.keep_with_next = True
        if content.title == "资金流图例":
            picture = doc.add_picture(BytesIO(graph_png(view, renderer=renderer)), width=Inches(7))
            picture._inline.docPr.set("descr", "资金流方向与中文差异状态 完整记录见明细附录")
        if content.headers:
            table = doc.add_table(rows=1, cols=len(content.headers))
            table.autofit = False
            widths = (
                (2.6, 1.3, 3.1)
                if content.title == "完整资金流明细"
                else (7 / len(content.headers),) * len(content.headers)
            )
            for column, width in zip(table.columns, widths, strict=True):
                column.width = Inches(width)
            for cell, title in zip(table.rows[0].cells, content.headers, strict=True):
                cell.text = title
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EEEEEE")
                cell._tc.get_or_add_tcPr().append(shade)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in content.rows:
                for cell, value in zip(table.add_row().cells, values, strict=True):
                    cell.text = value
            for row in table.rows:
                no_split = OxmlElement("w:cantSplit")
                row._tr.get_or_add_trPr().append(no_split)
                for cell, width in zip(row.cells, widths, strict=True):
                    cell.width = Inches(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    properties = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for edge in ("top", "left", "bottom", "right"):
                        border = OxmlElement(f"w:{edge}")
                        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                            border.set(qn(f"w:{key}"), value)
                        borders.append(border)
                    properties.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for edge in ("top", "left", "bottom", "right"):
                        margin = OxmlElement(f"w:{edge}")
                        margin.set(qn("w:w"), "100")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    properties.append(margins)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(4)
                        for run in paragraph.runs:
                            run.font.size = Pt(9.5 if content.title == "完整资金流明细" else 10)
    output = BytesIO()
    doc.save(output)
    return _embed_font(output.getvalue())
